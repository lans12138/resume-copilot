"""Bounded PDF/DOCX text extraction with stable source locators."""

from __future__ import annotations

import time
from collections.abc import Callable
from dataclasses import dataclass
from typing import BinaryIO, Protocol

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from backend.app.documents.validation import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE


@dataclass(frozen=True, slots=True)
class ParseLimits:
    max_pdf_pages: int
    max_extracted_chars: int
    timeout_seconds: int

    def __post_init__(self) -> None:
        if self.max_pdf_pages <= 0:
            raise ValueError("max_pdf_pages must be positive")
        if self.max_extracted_chars <= 0:
            raise ValueError("max_extracted_chars must be positive")
        if self.timeout_seconds <= 0:
            raise ValueError("timeout_seconds must be positive")


@dataclass(frozen=True, slots=True)
class PdfLocator:
    page_number: int
    block_index: int
    char_start: int
    char_end: int


@dataclass(frozen=True, slots=True)
class DocxParagraphLocator:
    paragraph_index: int
    char_start: int
    char_end: int


@dataclass(frozen=True, slots=True)
class DocxTableLocator:
    table_index: int
    row_index: int
    cell_index: int
    char_start: int
    char_end: int


DocumentLocator = PdfLocator | DocxParagraphLocator | DocxTableLocator


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    block_index: int
    text: str
    locator: DocumentLocator


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    media_type: str
    full_text: str
    blocks: tuple[ParsedBlock, ...]
    page_count: int | None = None
    paragraph_count: int | None = None
    table_count: int | None = None
    warnings: tuple[str, ...] = ()


class DocumentParseError(Exception):
    def __init__(self, code: str, safe_message: str, *, retryable: bool = False) -> None:
        super().__init__(safe_message)
        self.code = code
        self.safe_message = safe_message
        self.retryable = retryable


class DocumentParser(Protocol):
    media_types: frozenset[str]
    version: str

    def parse(self, stream: BinaryIO, limits: ParseLimits) -> ParsedDocument: ...


class _Deadline:
    def __init__(self, seconds: int, clock: Callable[[], float]) -> None:
        self._clock = clock
        self._expires_at = clock() + seconds

    def check(self) -> None:
        if self._clock() >= self._expires_at:
            raise DocumentParseError(
                "PARSER_TIMEOUT",
                "文档解析超时",
                retryable=True,
            )


class PyMuPdfParser:
    media_types = frozenset({PDF_MEDIA_TYPE})

    def __init__(self, *, version: str = "v1", clock: Callable[[], float] = time.monotonic) -> None:
        self.version = version
        self._clock = clock

    def parse(self, stream: BinaryIO, limits: ParseLimits) -> ParsedDocument:
        deadline = _Deadline(limits.timeout_seconds, self._clock)
        payload = _read_from_start(stream)
        deadline.check()
        try:
            document = pymupdf.open(  # type: ignore[no-untyped-call]
                stream=payload, filetype="pdf"
            )
        except (pymupdf.FileDataError, RuntimeError) as error:
            raise DocumentParseError("INVALID_PDF", "PDF 文件无法解析") from error

        try:
            if document.needs_pass:
                raise DocumentParseError("ENCRYPTED_PDF", "不支持加密 PDF")
            if document.page_count > limits.max_pdf_pages:
                raise DocumentParseError(
                    "PDF_PAGE_LIMIT_EXCEEDED",
                    "PDF 页数超过允许的上限",
                )

            blocks: list[ParsedBlock] = []
            character_count = 0
            for page_index in range(document.page_count):
                deadline.check()
                page = document.load_page(page_index)  # type: ignore[no-untyped-call]
                for source_block_index, raw_block in enumerate(
                    page.get_text("blocks", sort=True)
                ):
                    deadline.check()
                    if len(raw_block) > 7 and raw_block[7] != 0:
                        continue
                    text = _normalized_text(str(raw_block[4]))
                    if not text:
                        continue
                    character_count = _checked_character_count(
                        character_count,
                        text,
                        limits.max_extracted_chars,
                    )
                    blocks.append(
                        ParsedBlock(
                            block_index=len(blocks),
                            text=text,
                            locator=PdfLocator(
                                page_number=page_index + 1,
                                block_index=source_block_index,
                                char_start=0,
                                char_end=len(text),
                            ),
                        )
                    )
            return _parsed_document(
                PDF_MEDIA_TYPE,
                blocks,
                page_count=document.page_count,
            )
        finally:
            document.close()  # type: ignore[no-untyped-call]


class PythonDocxParser:
    media_types = frozenset({DOCX_MEDIA_TYPE})

    def __init__(self, *, version: str = "v1", clock: Callable[[], float] = time.monotonic) -> None:
        self.version = version
        self._clock = clock

    def parse(self, stream: BinaryIO, limits: ParseLimits) -> ParsedDocument:
        deadline = _Deadline(limits.timeout_seconds, self._clock)
        _rewind(stream)
        try:
            document = Document(stream)
        except Exception as error:
            raise DocumentParseError("INVALID_DOCX", "DOCX 文件无法解析") from error
        finally:
            _rewind(stream)
        deadline.check()

        blocks: list[ParsedBlock] = []
        character_count = 0
        paragraph_index = 0
        table_index = 0
        for content in document.iter_inner_content():
            deadline.check()
            if isinstance(content, Paragraph):
                text = _normalized_text(content.text)
                if text:
                    character_count = _checked_character_count(
                        character_count,
                        text,
                        limits.max_extracted_chars,
                    )
                    blocks.append(
                        ParsedBlock(
                            block_index=len(blocks),
                            text=text,
                            locator=DocxParagraphLocator(
                                paragraph_index=paragraph_index,
                                char_start=0,
                                char_end=len(text),
                            ),
                        )
                    )
                paragraph_index += 1
                continue

            if isinstance(content, Table):
                for row_index, row in enumerate(content.rows):
                    for cell_index, cell in enumerate(row.cells):
                        deadline.check()
                        text = _normalized_text(cell.text)
                        if not text:
                            continue
                        character_count = _checked_character_count(
                            character_count,
                            text,
                            limits.max_extracted_chars,
                        )
                        blocks.append(
                            ParsedBlock(
                                block_index=len(blocks),
                                text=text,
                                locator=DocxTableLocator(
                                    table_index=table_index,
                                    row_index=row_index,
                                    cell_index=cell_index,
                                    char_start=0,
                                    char_end=len(text),
                                ),
                            )
                        )
                table_index += 1
        return _parsed_document(
            DOCX_MEDIA_TYPE,
            blocks,
            paragraph_count=len(document.paragraphs),
            table_count=len(document.tables),
        )


def parser_for_media_type(media_type: str, *, version: str = "v1") -> DocumentParser:
    if media_type == PDF_MEDIA_TYPE:
        return PyMuPdfParser(version=version)
    if media_type == DOCX_MEDIA_TYPE:
        return PythonDocxParser(version=version)
    raise DocumentParseError("UNSUPPORTED_MEDIA", "没有适用于该文件类型的解析器")


def _parsed_document(
    media_type: str,
    blocks: list[ParsedBlock],
    *,
    page_count: int | None = None,
    paragraph_count: int | None = None,
    table_count: int | None = None,
) -> ParsedDocument:
    if not blocks:
        raise DocumentParseError(
            "EMPTY_TEXT",
            "未检测到可提取文本，文件可能是扫描件或空文档",
        )
    return ParsedDocument(
        media_type=media_type,
        full_text="\n\n".join(block.text for block in blocks),
        blocks=tuple(blocks),
        page_count=page_count,
        paragraph_count=paragraph_count,
        table_count=table_count,
    )


def _checked_character_count(current: int, text: str, maximum: int) -> int:
    updated = current + (2 if current else 0) + len(text)
    if updated > maximum:
        raise DocumentParseError(
            "EXTRACTED_TEXT_LIMIT_EXCEEDED",
            "提取文本长度超过允许的上限",
        )
    return updated


def _normalized_text(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def _read_from_start(stream: BinaryIO) -> bytes:
    _rewind(stream)
    try:
        return stream.read()
    finally:
        _rewind(stream)


def _rewind(stream: BinaryIO) -> None:
    stream.seek(0)
