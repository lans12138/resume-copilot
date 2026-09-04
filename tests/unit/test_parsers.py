"""PDF/DOCX parser limits, error taxonomy, and locator tests."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterator

import pymupdf
import pytest
from docx import Document

from backend.app.documents.parsers import (
    DocumentParseError,
    DocxParagraphLocator,
    DocxTableLocator,
    ParseLimits,
    PdfLocator,
    PyMuPdfParser,
    PythonDocxParser,
    parser_for_media_type,
)
from backend.app.documents.validation import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE


def limits(*, pages: int = 10, characters: int = 10_000, timeout: int = 5) -> ParseLimits:
    return ParseLimits(
        max_pdf_pages=pages,
        max_extracted_chars=characters,
        timeout_seconds=timeout,
    )


def make_pdf(*pages: str, password: str | None = None) -> io.BytesIO:
    document = pymupdf.open()  # type: ignore[no-untyped-call]
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    options: dict[str, object] = {}
    if password is not None:
        options = {
            "encryption": pymupdf.PDF_ENCRYPT_AES_256,  # type: ignore[attr-defined]
            "owner_pw": password,
            "user_pw": password,
        }
    payload = document.tobytes(**options)  # type: ignore[no-untyped-call]
    document.close()  # type: ignore[no-untyped-call]
    return io.BytesIO(payload)


def make_docx() -> io.BytesIO:
    document = Document()
    document.add_paragraph("Candidate Alice")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "FastAPI"
    document.add_paragraph("Python Engineer")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def make_malformed_docx() -> io.BytesIO:
    source = make_docx()
    malformed = io.BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(malformed, "w") as target:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename == "word/document.xml":
                payload = b"<w:document"
            target.writestr(info, payload)
    malformed.seek(0)
    return malformed


def error_code(error: pytest.ExceptionInfo[DocumentParseError]) -> str:
    return error.value.code


def test_pdf_parser_extracts_blocks_with_page_locators() -> None:
    result = PyMuPdfParser(version="test-v1").parse(
        make_pdf("Candidate Alice", "Python Engineer"),
        limits(),
    )

    assert result.media_type == PDF_MEDIA_TYPE
    assert result.page_count == 2
    assert result.full_text == "Candidate Alice\n\nPython Engineer"
    assert [block.block_index for block in result.blocks] == [0, 1]
    first = result.blocks[0]
    assert isinstance(first.locator, PdfLocator)
    assert first.locator.page_number == 1
    assert first.text[first.locator.char_start : first.locator.char_end] == "Candidate Alice"


def test_docx_parser_extracts_paragraph_and_table_locators() -> None:
    result = PythonDocxParser().parse(make_docx(), limits())

    assert result.media_type == DOCX_MEDIA_TYPE
    assert result.paragraph_count == 2
    assert result.table_count == 1
    assert [block.text for block in result.blocks] == [
        "Candidate Alice",
        "Skill",
        "FastAPI",
        "Python Engineer",
    ]
    paragraph = result.blocks[0]
    assert isinstance(paragraph.locator, DocxParagraphLocator)
    assert paragraph.locator.paragraph_index == 0
    table_cell = result.blocks[2]
    assert isinstance(table_cell.locator, DocxTableLocator)
    assert (table_cell.locator.table_index, table_cell.locator.row_index) == (0, 0)
    assert table_cell.locator.cell_index == 1
    assert table_cell.text[table_cell.locator.char_start : table_cell.locator.char_end] == "FastAPI"


@pytest.mark.parametrize(
    ("stream", "parser", "expected_code"),
    [
        (make_pdf(""), PyMuPdfParser(), "EMPTY_TEXT"),
        (io.BytesIO(b"not-pdf"), PyMuPdfParser(), "INVALID_PDF"),
        (io.BytesIO(b"not-docx"), PythonDocxParser(), "INVALID_DOCX"),
        (make_malformed_docx(), PythonDocxParser(), "INVALID_DOCX"),
    ],
)
def test_parser_classifies_empty_and_invalid_documents(
    stream: io.BytesIO,
    parser: PyMuPdfParser | PythonDocxParser,
    expected_code: str,
) -> None:
    with pytest.raises(DocumentParseError) as error:
        parser.parse(stream, limits())
    assert error_code(error) == expected_code
    assert not error.value.retryable


def test_pdf_parser_rejects_encryption_page_limit_and_character_limit() -> None:
    with pytest.raises(DocumentParseError) as encrypted:
        PyMuPdfParser().parse(make_pdf("secret", password="password"), limits())
    assert error_code(encrypted) == "ENCRYPTED_PDF"

    with pytest.raises(DocumentParseError) as pages:
        PyMuPdfParser().parse(make_pdf("one", "two"), limits(pages=1))
    assert error_code(pages) == "PDF_PAGE_LIMIT_EXCEEDED"

    with pytest.raises(DocumentParseError) as characters:
        PyMuPdfParser().parse(make_pdf("too many characters"), limits(characters=5))
    assert error_code(characters) == "EXTRACTED_TEXT_LIMIT_EXCEEDED"


def test_docx_parser_enforces_character_limit() -> None:
    with pytest.raises(DocumentParseError) as error:
        PythonDocxParser().parse(make_docx(), limits(characters=5))
    assert error_code(error) == "EXTRACTED_TEXT_LIMIT_EXCEEDED"


def test_character_limit_includes_block_separators() -> None:
    document = Document()
    document.add_paragraph("abc")
    document.add_paragraph("def")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)

    with pytest.raises(DocumentParseError) as error:
        PythonDocxParser().parse(stream, limits(characters=7))
    assert error_code(error) == "EXTRACTED_TEXT_LIMIT_EXCEEDED"


def test_parser_timeout_is_retryable() -> None:
    values: Iterator[float] = iter((0.0, 0.0, 10.0))
    parser = PyMuPdfParser(clock=lambda: next(values))
    with pytest.raises(DocumentParseError) as error:
        parser.parse(make_pdf("Candidate Alice"), limits(timeout=1))
    assert error_code(error) == "PARSER_TIMEOUT"
    assert error.value.retryable


def test_parser_registry_uses_media_type_and_rejects_unknown_type() -> None:
    assert isinstance(parser_for_media_type(PDF_MEDIA_TYPE), PyMuPdfParser)
    assert isinstance(parser_for_media_type(DOCX_MEDIA_TYPE), PythonDocxParser)
    with pytest.raises(DocumentParseError) as error:
        parser_for_media_type("text/plain")
    assert error_code(error) == "UNSUPPORTED_MEDIA"
